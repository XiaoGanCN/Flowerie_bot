import json
import re
import base64
import httpx
from io import BytesIO, StringIO
from typing import Tuple, Optional, List, Dict, Any
from loguru import logger

from src.config import Settings

# 可选依赖
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None
try:
    import docx
except ImportError:
    docx = None
try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    import csv
except ImportError:
    csv = None


class FileParser:
    def __init__(self, config: Settings):
        self.config = config
        self.file_cache: Dict[str, Tuple[str, float]] = {}  # url -> (content, timestamp)
        self._client: Optional[httpx.AsyncClient] = None  # 复用的 HTTP 客户端（懒创建）

    def _get_client(self, timeout: float = 30) -> httpx.AsyncClient:
        """复用一个 AsyncClient，避免每次请求都新建连接（省连接开销）"""
        if self._client is None or self._client.is_closed:
            self._client = httpx.AsyncClient(timeout=timeout)
        return self._client

    def _cap(self, text: str) -> str:
        """P1-4 资源上限：提取文本统一截断到 MAX_FILE_TEXT_CHARS，防止超长内容烧 token。"""
        if not text:
            return ""
        limit = getattr(self.config, "MAX_FILE_TEXT_CHARS", 8000)
        if len(text) > limit:
            return text[:limit] + "\n...(内容过长已截断)"
        return text

    # ========== 新增：通过 NapCat HTTP API 获取并解析文件 ==========
    async def fetch_and_parse_file(self, file_id: str, file_name: str) -> Tuple[str, bool]:
        """
        调用 NapCat /get_file 接口获取文件内容，并调用 decode_napcat_file_response 解析
        返回: (提取的文本内容, 是否成功)
        """
        if not file_id:
            return "", False

        try:
            client = self._get_client(timeout=30)
            resp = await client.get(
                f"{self.config.HTTP_API_BASE}/get_file",
                params={"file_id": file_id},
            )
            if resp.status_code != 200:
                logger.error(f"Fetch file {file_id} failed: HTTP {resp.status_code}")
                return "", False

            # 调用已有的解码方法
            return self.decode_napcat_file_response(resp.text, file_name)

        except httpx.TimeoutException:
            logger.error(f"Fetch file {file_id} timeout")
            return "", False
        except httpx.HTTPError as e:
            logger.error(f"Fetch file {file_id} HTTP error: {e}")
            return "", False
        except Exception as e:
            logger.exception(f"Fetch file {file_id} unexpected error: {e}")
            return "", False

    # ========== 解码 NapCat 文件响应 ==========
    def decode_napcat_file_response(self, response_text: str, file_name: str) -> Tuple[str, bool]:
        """从 /get_file 返回的 JSON 中提取 base64 并解码"""
        try:
            data = json.loads(response_text)
            if data.get("retcode") != 0:
                return "", False
            file_data = data.get("data", {})
            b64 = file_data.get("base64", "")
            if not b64:
                return "", False
            content_bytes = base64.b64decode(b64)
            ext = file_name.split('.')[-1].lower() if '.' in file_name else ''
            extracted_text = ""

            if ext == 'txt' or ext == '':
                for enc in ['utf-8', 'gbk']:
                    try:
                        extracted_text = content_bytes.decode(enc)
                        return extracted_text, True
                    except UnicodeDecodeError:
                        continue
                return content_bytes.decode('utf-8', errors='ignore'), True

            elif ext == 'pdf':
                if PyPDF2 is None:
                    return "", False
                try:
                    pdf_reader = PyPDF2.PdfReader(BytesIO(content_bytes))
                    pages = [page.extract_text() or "" for page in pdf_reader.pages[:self.config.MAX_PDF_PAGES]]
                    extracted_text = self._cap("\n".join(pages))
                    return extracted_text, True
                except Exception as e:
                    logger.error(f"PDF parse error: {e}")
                    return "", False

            elif ext == 'docx':
                if docx is None:
                    return "", False
                try:
                    doc = docx.Document(BytesIO(content_bytes))
                    paragraphs = [p.text for p in doc.paragraphs]
                    extracted_text = self._cap("\n".join(paragraphs))
                    return extracted_text, True
                except Exception as e:
                    logger.error(f"DOCX parse error: {e}")
                    return "", False

            elif ext == 'xlsx':
                if openpyxl is None:
                    return "", False
                try:
                    wb = openpyxl.load_workbook(BytesIO(content_bytes), data_only=True)
                    rows = []
                    cell_count = 0
                    for sheet in wb.worksheets:
                        rows.append(f"--- 工作表: {sheet.title} ---")
                        for row in sheet.iter_rows(values=True):
                            rows.append("\t".join([str(cell) if cell is not None else "" for cell in row]))
                            cell_count += len(row)
                            if cell_count >= self.config.MAX_EXCEL_CELLS:
                                rows.append("...(已达到解析上限)")
                                break
                        if cell_count >= self.config.MAX_EXCEL_CELLS:
                            break
                    extracted_text = self._cap("\n".join(rows))
                    return extracted_text, True
                except Exception as e:
                    logger.error(f"XLSX parse error: {e}")
                    return "", False

            elif ext == 'csv':
                if csv is None or StringIO is None:
                    return "", False
                try:
                    for enc in ['utf-8', 'gbk', 'gb2312']:
                        try:
                            text_content = content_bytes.decode(enc)
                            break
                        except UnicodeDecodeError:
                            continue
                    else:
                        text_content = content_bytes.decode('utf-8', errors='ignore')
                    reader = csv.reader(StringIO(text_content))
                    rows = [",".join(row) for row in reader][:self.config.MAX_CSV_ROWS]
                    extracted_text = self._cap("\n".join(rows))
                    return extracted_text, True
                except Exception as e:
                    logger.error(f"CSV parse error: {e}")
                    return "", False
            else:
                # 尝试作为文本
                try:
                    extracted_text = self._cap(content_bytes.decode('utf-8'))
                    return extracted_text, True
                except:
                    return "", False
        except json.JSONDecodeError:
            return response_text, True
        except Exception as e:
            logger.error(f"File decode error: {e}")
            return "", False

    # ========== 提取合并转发消息 ==========
    async def extract_forward_messages(self, message_array: List[Dict]) -> Tuple[str, List[str], bool]:
        """提取合并转发消息中的文本与图片 URL。

        返回: (转发文本, 转发内的图片url列表, 是否有转发)
        图片 url 由调用方逐个交给视觉模型识图（让花璃看到转发里的每一张图）。
        """
        import httpx

        def extract_all_text(obj, sender="未知", prefix="", urls=None):
            results = []
            if isinstance(obj, dict):
                sender_val = obj.get("sender", {})
                if isinstance(sender_val, dict):
                    sender = sender_val.get("user_id", sender)
                elif isinstance(sender_val, (int, str)):
                    sender = str(sender_val)
                # 收集 image 段的 url（NapCat 合并转发里的图片同样带 url）
                if obj.get("type") == "image":
                    img_url = (obj.get("data") or {}).get("url", "")
                    if img_url and urls is not None and img_url not in urls:
                        urls.append(img_url)
                if "text" in obj and isinstance(obj["text"], str):
                    if obj["text"].strip():
                        results.append(f"[用户{sender}]：{obj['text']}")
                for key, value in obj.items():
                    if key == "text":
                        continue
                    results.extend(extract_all_text(value, sender, prefix + key + ".", urls))
            elif isinstance(obj, list):
                for idx, item in enumerate(obj):
                    results.extend(extract_all_text(item, sender, prefix + f"[{idx}]", urls))
            return results

        async def fetch_forward_messages(forward_id) -> Optional[List[Dict]]:
            """通过 NapCat HTTP API 拉取一次转发内容。"""
            try:
                client = self._get_client(timeout=10)
                resp = await client.get(
                    f"{self.config.HTTP_API_BASE}/get_forward_msg",
                    params={"message_id": forward_id}
                )
                if resp.status_code == 200:
                    result = resp.json()
                    if result.get("retcode") == 0:
                        return result.get("data", {}).get("messages")
            except Exception as e:
                logger.error(f"Get forward msg error: {e}")
            return None

        async def resolve_nested_forwards(node, depth: int = 0):
            """递归展开嵌套转发：把 forward 节点（无内联 messages 但有 id）拉取后展开。

            这样聊天记录里嵌套的聊天记录、以及转发里的转发都能解析出来。
            深度上限 MAX_FORWARD_DEPTH（P1-5），防止恶意多层嵌套拖垮服务。
            """
            if depth > getattr(self.config, "MAX_FORWARD_DEPTH", 5):
                return node
            if isinstance(node, list):
                return [await resolve_nested_forwards(item, depth) for item in node]
            if not isinstance(node, dict):
                return node
            if node.get("type") == "forward":
                data = node.get("data") or {}
                messages = data.get("messages")
                if not messages:
                    fid = data.get("id")
                    if fid:
                        fetched = await fetch_forward_messages(fid)
                        if fetched:
                            data["messages"] = fetched
                            messages = fetched
                if messages:
                    return await resolve_nested_forwards(messages, depth + 1)
                return node
            # 普通节点：递归展开所有字段（含 message 数组里的嵌套 forward）
            return {key: await resolve_nested_forwards(value, depth) for key, value in node.items()}

        for msg in message_array:
            if msg.get("type") == "forward":
                forward_data = msg.get("data", {})
                messages = forward_data.get("messages")
                if not messages:
                    messages = await fetch_forward_messages(forward_data.get("id"))
                    if not messages:
                        continue
                # 先展开嵌套转发，再一次性提取文本与图片 url
                resolved = await resolve_nested_forwards(messages)
                image_urls: List[str] = []
                text_lines = extract_all_text(resolved, urls=image_urls)
                if text_lines or image_urls:
                    return "\n".join(text_lines), image_urls, True
                return "", image_urls, False
        return "", [], False

    # ========== 提取 JSON 卡片内容 ==========
    def extract_json_card_content(self, message_array: List[Dict]) -> Tuple[str, bool]:
        """递归提取 JSON 卡片中的所有文本"""
        def collect_strings(obj, collected: set):
            if isinstance(obj, dict):
                for key, value in obj.items():
                    if key in ['url', 'jumpUrl', 'preview', 'icon', 'appid', 'uin', 'scene', 'token', 'ctime', 'width', 'height', 'forward', 'autoSize']:
                        continue
                    collect_strings(value, collected)
            elif isinstance(obj, list):
                for item in obj:
                    collect_strings(item, collected)
            elif isinstance(obj, str):
                if len(obj.strip()) > 1 and not obj.strip().isdigit():
                    collected.add(obj.strip())

        for msg in message_array:
            if msg.get("type") != "json":
                continue
            data = msg.get("data", {})
            json_str = data.get("data") or data.get("content") or data.get("text")
            if isinstance(json_str, dict):
                card_data = json_str
            elif isinstance(json_str, str):
                try:
                    card_data = json.loads(json_str)
                except json.JSONDecodeError:
                    continue
            else:
                if isinstance(data, dict):
                    card_data = data
                else:
                    continue
            collected = set()
            collect_strings(card_data, collected)
            if collected:
                card_text = "；".join(list(collected))
                return f"卡片内容：{card_text}", True
        return "", False

    # ========== 提取 @ 和纯文本 ==========
    def extract_mention_and_text(self, message_array: List[Dict], bot_qq: int) -> Tuple[str, bool]:
        """提取纯文本和是否@机器人"""
        self_id = str(bot_qq)
        is_mentioned = False
        text_parts = []
        for msg in message_array:
            if msg.get("type") == "at":
                qq = str(msg.get("data", {}).get("qq", ""))
                if qq == self_id:
                    is_mentioned = True
            elif msg.get("type") == "text":
                text_parts.append(msg.get("data", {}).get("text", ""))
        return "".join(text_parts).strip(), is_mentioned